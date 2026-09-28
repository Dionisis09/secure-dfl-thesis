from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Secure_DFL_Defense_Guide_GR.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(10)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, bold=True, color="0B2545")
        set_cell_shading(hdr[i], "E8EEF5")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value))
    set_table_widths(table, widths)
    doc.add_paragraph()
    return table


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "F4F6F9")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor(31, 77, 120)
    r.font.size = Pt(10.5)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    r2.font.size = Pt(10.5)
    table.autofit = False
    cell.width = Inches(6.5)
    doc.add_paragraph()


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbers(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Consolas"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    r.font.size = Pt(9)


def add_qa(doc, question, answer):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    q = p.add_run("Ερώτηση: ")
    q.bold = True
    q.font.color.rgb = RGBColor(31, 77, 120)
    p.add_run(question)
    p2 = doc.add_paragraph()
    p2.paragraph_format.left_indent = Inches(0.25)
    p2.paragraph_format.space_after = Pt(6)
    a = p2.add_run("Απάντηση: ")
    a.bold = True
    p2.add_run(answer)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Title", 22, "0B2545", 0, 10),
        ("Subtitle", 12, "555555", 0, 8),
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        if name != "Subtitle":
            style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    for name in ["List Bullet", "List Number"]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.167

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Secure DFL - Οδηγός Κατανόησης και Άμυνας Διπλωματικής")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(85, 85, 85)


def build():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Secure Decentralized Federated Learning")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Μεγάλος Οδηγός Κατανόησης, Παρουσίασης και Άμυνας Διπλωματικής")

    meta = [
        ["Σκοπός αρχείου", "Να μπορείς να εξηγήσεις από το μηδέν τι υλοποιήθηκε, γιατί υλοποιήθηκε και πώς να απαντήσεις σε ερωτήσεις."],
        ["Πεδίο", "Decentralized Federated Learning, privacy mechanisms, HE, adaptive hybrid security, blockchain audit, platform demo."],
        ["Τρόπος χρήσης", "Διάβασέ το σαν oral defense guide. Δεν αντικαθιστά την τεχνική αναφορά, τη συμπληρώνει."],
    ]
    add_table(doc, ["Πεδίο", "Περιγραφή"], meta, [1.6, 4.9])

    add_callout(
        doc,
        "Κεντρική ιδέα σε μία πρόταση",
        "Η εργασία υλοποιεί ένα decentralized federated learning σύστημα όπου οι clients κρατούν τα δεδομένα τοπικά, "
        "εκπαιδεύουν μοντέλα, ανταλλάσσουν updates με γείτονες και χρησιμοποιούν διαφορετικούς μηχανισμούς προστασίας "
        "και audit για να μετρηθεί το trade-off ανάμεσα σε accuracy, privacy, χρόνο και communication cost.",
    )

    doc.add_heading("Περιεχόμενα ανάγνωσης", level=1)
    add_numbers(doc, [
        "Τι πρόβλημα λύνει η εργασία.",
        "Τι είναι Federated Learning και τι αλλάζει στο Decentralized Federated Learning.",
        "Πώς δουλεύει το training pipeline ανά round.",
        "Τι σημαίνουν IID και non-IID δεδομένα.",
        "Ποιος είναι ο ρόλος των datasets και των topologies.",
        "Τι κάνει κάθε security mode: none, masking, pairwise masking, selective HE, adaptive hybrid.",
        "Πώς ερμηνεύονται τα αποτελέσματα και τα plots.",
        "Τι ρόλο έχει το blockchain/audit layer.",
        "Τι δείχνει το Docker/platform demo.",
        "Πιθανές ερωτήσεις καθηγητή και προτεινόμενες απαντήσεις.",
    ])

    doc.add_page_break()

    doc.add_heading("1. Το πρόβλημα που λύνει η εργασία", level=1)
    doc.add_paragraph(
        "Σε πολλά σύγχρονα συστήματα machine learning τα δεδομένα βρίσκονται κατανεμημένα σε διαφορετικές συσκευές ή οργανισμούς. "
        "Η κλασική λύση θα ήταν να συγκεντρωθούν όλα τα δεδομένα σε έναν server και να εκπαιδευτεί εκεί ένα μοντέλο. "
        "Αυτό όμως συχνά δεν είναι αποδεκτό λόγω privacy, κανονιστικών περιορισμών, latency, κόστους ή έλλειψης εμπιστοσύνης σε έναν κεντρικό φορέα."
    )
    doc.add_paragraph(
        "Η εργασία προσεγγίζει αυτό το πρόβλημα με Decentralized Federated Learning. Κάθε client κρατάει τα δεδομένα του τοπικά, "
        "εκπαιδεύει το δικό του μοντέλο και ανταλλάσσει μόνο model updates ή model parameters με γειτονικούς clients. "
        "Πάνω σε αυτό προστίθενται μηχανισμοί προστασίας ώστε η ανταλλαγή των updates να μην γίνεται γυμνά."
    )
    add_callout(
        doc,
        "Τι να πεις αν σε ρωτήσουν γιατί έχει αξία",
        "Η αξία δεν είναι μόνο ότι τρέχει ένα FL πείραμα. Η αξία είναι ότι υπάρχει πλήρης πειραματικό pipeline με decentralized topologies, "
        "πολλαπλά security modes, metrics, plots, blockchain-style audit και containerized demo. Άρα η εργασία δεν είναι μόνο θεωρητική.",
    )

    doc.add_heading("2. Federated Learning από το μηδέν", level=1)
    doc.add_paragraph(
        "Στο Federated Learning το μοντέλο μαθαίνει από δεδομένα που παραμένουν στους clients. Οι clients δεν στέλνουν τα raw data. "
        "Στέλνουν ενημερώσεις του μοντέλου, όπως weights, gradients ή state_dict parameters. Αυτές οι ενημερώσεις συνδυάζονται ώστε να προκύψει ένα καλύτερο κοινό μοντέλο."
    )
    add_table(doc, ["Κλασικό ML", "Federated Learning"], [
        ["Τα δεδομένα συγκεντρώνονται κεντρικά.", "Τα δεδομένα μένουν τοπικά στους clients."],
        ["Ο server εκπαιδεύει το μοντέλο.", "Κάθε client εκπαιδεύει τοπικά."],
        ["Υψηλή εξάρτηση από κεντρικό dataset.", "Εκπαίδευση πάνω σε κατανεμημένα δεδομένα."],
        ["Μεγαλύτερος privacy κίνδυνος από data centralization.", "Μειώνεται η ανάγκη μεταφοράς raw data."],
    ], [3.25, 3.25])

    doc.add_heading("3. Τι είναι Decentralized Federated Learning", level=1)
    doc.add_paragraph(
        "Στο κλασικό FL υπάρχει συνήθως κεντρικός server που κάνει aggregation. Στο Decentralized Federated Learning δεν υπάρχει κεντρικός aggregation server. "
        "Οι clients επικοινωνούν μεταξύ τους σύμφωνα με έναν γράφο επικοινωνίας, δηλαδή μία topology. Κάθε client λαμβάνει updates από τους γείτονές του και κάνει τοπικό aggregation."
    )
    add_code(doc, "Client trains locally -> protects update -> sends to neighbors -> receives neighbor updates -> aggregates -> next round")
    add_callout(
        doc,
        "Σημαντική φράση άμυνας",
        "Το decentralized setup μειώνει την εξάρτηση από έναν κεντρικό server και ταιριάζει περισσότερο σε peer-to-peer, edge και IoT περιβάλλοντα.",
    )

    doc.add_heading("4. Ο κύκλος ενός communication round", level=1)
    add_numbers(doc, [
        "Ορίζονται clients, dataset split, topology και security mode.",
        "Κάθε client παίρνει το δικό του local subset δεδομένων.",
        "Κάθε client εκπαιδεύει τοπικά το μοντέλο του.",
        "Το παραγόμενο model update προστατεύεται ανάλογα με το security mode.",
        "Ο client στέλνει το προστατευμένο update στους neighbors.",
        "Ο client λαμβάνει updates από τους neighbors.",
        "Γίνεται aggregation με averaging ώστε να ενημερωθεί το τοπικό μοντέλο.",
        "Καταγράφονται metrics: accuracy, loss, time, communication bytes, security overhead.",
        "Προαιρετικά δημιουργείται audit/blockchain record για τον round.",
    ])

    doc.add_heading("5. Τι σημαίνει state_dict, update και aggregation", level=1)
    doc.add_paragraph(
        "Στην PyTorch το state_dict είναι μία δομή που περιέχει όλα τα trainable parameters του μοντέλου. "
        "Για ένα neural network περιλαμβάνει weights και biases ανά layer. Όταν οι clients επικοινωνούν, αυτό που πρακτικά συγκρίνεται ή μοιράζεται είναι αυτή η κατάσταση του μοντέλου ή η διαφορά της από προηγούμενη κατάσταση."
    )
    add_table(doc, ["Όρος", "Πρακτική σημασία"], [
        ["state_dict", "Το σύνολο των weights/biases του μοντέλου σε συγκεκριμένη στιγμή."],
        ["model update", "Η αλλαγή του μοντέλου μετά την τοπική εκπαίδευση ή το νέο state που θα μοιραστεί."],
        ["aggregation", "Ο συνδυασμός πολλών updates, συνήθως με μέσο όρο."],
        ["communication round", "Ένας κύκλος local training, exchange και aggregation."],
    ], [1.7, 4.8])

    doc.add_heading("6. IID και non-IID δεδομένα", level=1)
    doc.add_paragraph(
        "Το IID σημαίνει Independent and Identically Distributed. Σε IID split οι clients έχουν παρόμοια κατανομή δεδομένων. "
        "Σε non-IID split οι clients έχουν διαφορετικές κατανομές. Αυτό είναι πιο ρεαλιστικό και πιο δύσκολο."
    )
    add_table(doc, ["IID", "non-IID"], [
        ["Οι clients έχουν παρόμοια classes.", "Οι clients έχουν biased ή άνισες classes."],
        ["Πιο εύκολη σύγκλιση.", "Πιο δύσκολη σύγκλιση."],
        ["Συνήθως υψηλότερη accuracy.", "Συνήθως χαμηλότερη accuracy."],
        ["Καλό sanity check.", "Πιο κοντά σε πραγματικά distributed δεδομένα."],
    ], [3.25, 3.25])
    add_callout(
        doc,
        "Πώς εξηγείς το χαμηλότερο non-IID accuracy",
        "Στο non-IID κάθε client βλέπει διαφορετικό κομμάτι της πραγματικότητας. Άρα τα local updates είναι πιο ανομοιογενή και το aggregation δυσκολεύεται περισσότερο.",
    )

    doc.add_heading("7. Datasets", level=1)
    add_table(doc, ["Dataset", "Γιατί χρησιμοποιείται", "Τι πρέπει να πεις"], [
        ["MNIST", "Γρήγορο και κλασικό benchmark.", "Χρησιμοποιείται για σταθερό baseline και έλεγχο σωστής λειτουργίας."],
        ["Fashion-MNIST", "Πιο δύσκολο από MNIST με ίδια περίπου input μορφή.", "Δείχνει ότι το pipeline δεν είναι δεμένο μόνο με digit classification."],
        ["CIFAR-10", "Πιο ρεαλιστικό και δύσκολο RGB dataset.", "Χρειάζεται ισχυρότερο μοντέλο, π.χ. CNN, για υψηλή accuracy."],
    ], [1.35, 2.5, 2.65])

    doc.add_heading("8. Topologies", level=1)
    add_table(doc, ["Topology", "Περιγραφή", "Trade-off"], [
        ["ring", "Κάθε client μιλάει με δύο γείτονες.", "Χαμηλό κόστος, πιο αργή διάδοση πληροφορίας."],
        ["fully_connected", "Κάθε client μιλάει με όλους.", "Γρήγορη διάδοση, μεγάλο communication cost."],
        ["star", "Ένας hub συνδέεται με όλους.", "Απλό, αλλά ο hub είναι κρίσιμο σημείο."],
        ["random", "Τυχαίος ελεγχόμενος peer-to-peer γράφος.", "Πιο ρεαλιστικό αλλά εξαρτάται από connectivity."],
        ["small_world", "Ring με shortcuts.", "Καλός συμβιβασμός κόστους και διάδοσης."],
    ], [1.45, 2.7, 2.35])

    doc.add_heading("9. Security mode: none", level=1)
    doc.add_paragraph(
        "Το none είναι το καθαρό baseline. Δεν εφαρμόζεται masking ή encryption. Χρησιμοποιείται ως σημείο σύγκρισης για όλα τα υπόλοιπα security modes."
    )
    add_callout(doc, "Τι να πεις", "Χωρίς baseline δεν μπορούμε να μετρήσουμε το κόστος που προσθέτει κάθε security mechanism.")

    doc.add_heading("10. Security mode: simple masking", level=1)
    doc.add_paragraph(
        "Στο simple masking δημιουργείται ένα τυχαίο mask με ίδια σχήματα tensor όπως τα parameters του μοντέλου. "
        "Το mask προστίθεται στο update πριν από τη διαμοίραση."
    )
    add_code(doc, "masked_update = update + mask")
    doc.add_paragraph(
        "Αυτή η έκδοση είναι απλοποιημένη simulation protected sharing. Είναι χρήσιμη για να δείξει πού μπαίνει το privacy layer, "
        "πώς μετριέται το masking overhead και ότι το pipeline συνεχίζει να δουλεύει."
    )
    add_callout(
        doc,
        "Προσοχή στην άμυνα",
        "Μην το παρουσιάσεις ως πλήρες secure aggregation protocol. Πες ότι είναι απλή controlled masking έκδοση και ότι το ισχυρότερο masking μέρος είναι το pairwise masking.",
    )

    doc.add_heading("11. Security mode: pairwise masking", level=1)
    doc.add_paragraph(
        "Το pairwise masking είναι πιο κοντά σε secure aggregation λογική. Κάθε ζευγάρι clients δημιουργεί masks που ακυρώνονται μεταξύ τους στο aggregation."
    )
    add_code(doc, "(update_A + m) + (update_B - m) = update_A + update_B")
    doc.add_paragraph(
        "Έτσι τα individual updates δεν εμφανίζονται raw κατά τη μεταφορά, αλλά το τελικό aggregated αποτέλεσμα παραμένει μαθηματικά σωστό. "
        "Η μικρή διαφορά που μπορεί να παρατηρηθεί είναι floating point cancellation error και όχι ουσιαστική αλλοίωση του training."
    )

    doc.add_heading("12. Security mode: selective Homomorphic Encryption", level=1)
    doc.add_paragraph(
        "Το Homomorphic Encryption επιτρέπει υπολογισμούς πάνω σε encrypted τιμές. Στο project χρησιμοποιείται selective HE με CKKS/TenSEAL λογική, "
        "δηλαδή δεν κρυπτογραφείται απαραίτητα όλο το μοντέλο, αλλά επιλεγμένα μέρη ή επιλεγμένοι στόχοι."
    )
    doc.add_paragraph(
        "Ο λόγος που δεν χρησιμοποιείται HE παντού είναι το κόστος. Τα ciphertexts είναι μεγαλύτερα από τα απλά tensors και η κρυπτογράφηση/αποκρυπτογράφηση προσθέτει χρόνο."
    )
    add_callout(
        doc,
        "Καθαρή απάντηση",
        "Το HE προσφέρει ισχυρότερη προστασία, αλλά έχει υψηλό overhead. Για αυτό χρησιμοποιείται επιλεκτικά και όχι αδιακρίτως.",
    )

    doc.add_heading("13. Security mode: adaptive hybrid", level=1)
    doc.add_paragraph(
        "Το adaptive hybrid είναι η βασική τεχνική συνεισφορά. Συνδυάζει pairwise masking και selective HE. "
        "Για κάθε round/client υπολογίζεται risk score. Οι πιο risky clients ή targets προστατεύονται με HE, ενώ οι υπόλοιποι χρησιμοποιούν pairwise masking."
    )
    add_code(doc, "risk score -> επιλογή HE targets -> HE για high-risk clients -> pairwise masking για τους υπόλοιπους")
    doc.add_paragraph(
        "Στο βασικό setup με 5 clients και HE ratio 0.40, αυτό σημαίνει περίπου 2 clients με HE ανά round και οι υπόλοιποι με pairwise masking. "
        "Έτσι το σύστημα προσπαθεί να πάρει περισσότερη προστασία από pure masking με μικρότερο κόστος από pure HE."
    )
    add_table(doc, ["Μέθοδος", "Προστασία", "Κόστος", "Σχόλιο"], [
        ["none", "Καμία", "Χαμηλό", "Χρειάζεται ως baseline."],
        ["pairwise masking", "Μεσαία", "Χαμηλό/μεσαίο", "Καλή ακύρωση masks."],
        ["selective HE", "Υψηλότερη", "Υψηλότερο", "Ακριβό λόγω ciphertexts."],
        ["adaptive hybrid", "Στοχευμένη", "Ελεγχόμενο", "Κύρια συνεισφορά/trade-off."],
    ], [1.45, 1.45, 1.35, 2.25])

    doc.add_heading("14. Γιατί η accuracy μένει κοντά ίδια", level=1)
    doc.add_paragraph(
        "Οι security μηχανισμοί δεν σχεδιάστηκαν για να αυξήσουν accuracy. Σχεδιάστηκαν για να προστατεύσουν την ανταλλαγή updates. "
        "Άρα όταν η accuracy παραμένει κοντά στο baseline, αυτό είναι θετικό αποτέλεσμα: δείχνει ότι η προστασία δεν κατέστρεψε το training."
    )
    add_callout(
        doc,
        "Αν φαίνεται μόνο μία γραμμή σε plot",
        "Δεν σημαίνει ότι λείπουν μέθοδοι. Συχνά οι καμπύλες επικαλύπτονται επειδή οι διαφορές στο loss είναι πάρα πολύ μικρές.",
    )

    doc.add_heading("15. Πώς διαβάζεις τα αποτελέσματα", level=1)
    add_bullets(doc, [
        "Το accuracy δείχνει πόσο σωστά ταξινομεί το μοντέλο στο test set.",
        "Το loss δείχνει πόσο λάθος είναι οι προβλέψεις, με χαμηλότερο loss να είναι καλύτερο.",
        "Το communication bytes δείχνει το κόστος μεταφοράς updates/metadata.",
        "Το masking time δείχνει τον χρόνο δημιουργίας/εφαρμογής masks.",
        "Το HE time δείχνει το κόστος κρυπτογράφησης/αποκρυπτογράφησης.",
        "Το blockchain verification δείχνει αν το audit chain παρέμεινε έγκυρο.",
    ])
    add_table(doc, ["Παρατήρηση", "Ερμηνεία"], [
        ["IID accuracy υψηλότερο από non-IID.", "Αναμενόμενο, γιατί το IID είναι πιο εύκολο."],
        ["Security modes έχουν παρόμοια accuracy.", "Η προστασία δεν αλλοιώνει ουσιαστικά το learning."],
        ["HE αυξάνει communication/time.", "Αναμενόμενο λόγω ciphertext expansion."],
        ["Adaptive hybrid μικρότερο κόστος από HE για όλους.", "Επιβεβαιώνει το trade-off της μεθόδου."],
    ], [2.4, 4.1])

    doc.add_heading("16. Blockchain και audit layer", level=1)
    doc.add_paragraph(
        "Το blockchain component είναι lightweight και εκπαιδευτικό. Δεν είναι production blockchain. Δεν έχει consensus, mining ή validator network. "
        "Χρησιμοποιείται για tamper-evident audit trail των communication rounds."
    )
    add_bullets(doc, [
        "Κάθε block έχει hash.",
        "Κάθε block δείχνει στο previous hash.",
        "Αν αλλάξει ένα block, χαλάει η αλυσίδα.",
        "Υπάρχει verification report.",
        "Υπάρχει tampering demonstration.",
        "Προστέθηκαν metadata, signatures και verification statuses.",
    ])
    add_callout(
        doc,
        "Σωστή διατύπωση",
        "Το blockchain δεν προστατεύει privacy. Προσφέρει integrity, traceability και auditability.",
    )

    doc.add_heading("17. Digital signatures στο audit", level=1)
    doc.add_paragraph(
        "Οι simulated ή real-key signatures βοηθούν ώστε κάθε update/block να μπορεί να συνδεθεί με συγκεκριμένο client identity. "
        "Στην εκπαιδευτική blockchain έκδοση οι signatures είναι deterministic simulation. Στο platform demo υπάρχουν πραγματικά generated keys για πιο ρεαλιστική παρουσίαση."
    )
    add_table(doc, ["Status", "Σημασία"], [
        ["VALID", "Hash chain και signatures είναι σωστά."],
        ["INVALID_SIGNATURE", "Το περιεχόμενο ή η υπογραφή δεν ταιριάζει με τον client."],
        ["INVALID_HASH_CHAIN", "Το previous hash linking έχει σπάσει."],
    ], [1.8, 4.7])

    doc.add_heading("18. Platform και Docker demo", level=1)
    doc.add_paragraph(
        "Το platform/Docker demo δείχνει ότι το σύστημα μπορεί να παρουσιαστεί και ως μικρή distributed εφαρμογή, όχι μόνο ως offline experiment script. "
        "Περιλαμβάνει nodes, signed messages, audit logs, dashboard και verifier."
    )
    add_bullets(doc, [
        "Δεν αντικαθιστά τα πειράματα της διπλωματικής.",
        "Δείχνει operational realism.",
        "Βοηθά στην παρουσίαση επειδή φαίνεται node status και audit status.",
        "Είναι proof-of-concept για πιθανή μελλοντική productization.",
    ])

    doc.add_heading("19. Εξήγηση ανά βασικό module κώδικα", level=1)
    doc.add_paragraph(
        "Αυτή η ενότητα είναι χρήσιμη αν ο καθηγητής σε ρωτήσει πού βρίσκεται κάθε λειτουργία στον κώδικα. "
        "Δεν χρειάζεται να θυμάσαι κάθε γραμμή, αλλά πρέπει να ξέρεις τον ρόλο κάθε αρχείου."
    )
    add_table(doc, ["Module / αρχείο", "Ρόλος στην εργασία", "Πώς το εξηγείς προφορικά"], [
        ["main.py", "Κεντρικός experiment runner.", "Εκεί ξεκινά το πείραμα, διαβάζονται τα argparse options και συντονίζονται clients, rounds, metrics και plots."],
        ["config.py", "Βασικές ρυθμίσεις.", "Κρατά default παραμέτρους ώστε τα experiments να είναι πιο οργανωμένα."],
        ["data_utils.py", "Φόρτωση dataset και split.", "Εκεί υλοποιείται το IID/non-IID partitioning των δεδομένων στους clients."],
        ["model.py", "Ορισμός neural network.", "Περιέχει το μοντέλο ταξινόμησης και υποστηρίζει διαφορετικά input shapes για MNIST/FMNIST/CIFAR-10."],
        ["topology.py", "Γράφοι επικοινωνίας.", "Ορίζει ποιοι clients είναι neighbors σε ring, fully connected, star, random και small-world."],
        ["masking.py", "Mask generation/application.", "Περιέχει simple masking και pairwise masking λογική."],
        ["he_utils.py", "Homomorphic encryption helpers.", "Περιέχει τις βοηθητικές λειτουργίες για selective HE με CKKS/TenSEAL λογική."],
        ["adaptive.py", "Risk score και adaptive selection.", "Υπολογίζει ποιοι clients/targets θα πάρουν HE αντί για masking."],
        ["blockchain.py", "Audit/hash chain.", "Δημιουργεί blocks, hashes, signatures και verification reports."],
        ["compare_results.py", "Συγκρίσεις και plots.", "Παίρνει CSV αποτελέσματα και παράγει συγκριτικά figures."],
        ["run_validation_experiments.py", "Validation experiments.", "Τρέχει πιο συστηματικά πειράματα για να μην βασιζόμαστε σε ένα μόνο run."],
        ["run_adaptive_sensitivity.py", "Sensitivity analysis.", "Μετρά τι γίνεται όταν αλλάζει το adaptive HE ratio."],
    ], [1.65, 2.25, 2.6])

    doc.add_heading("20. Πώς να παρουσιάσεις τα slides", level=1)
    doc.add_paragraph(
        "Αν δεν θυμάσαι ακριβώς κάθε slide, χρησιμοποίησε την παρακάτω λογική. Η παρουσίαση πρέπει να ακούγεται σαν ιστορία: "
        "πρόβλημα, λύση, baseline, security, αποτελέσματα, audit, συμπέρασμα."
    )
    add_table(doc, ["Τμήμα παρουσίασης", "Τι λες", "Τι αποφεύγεις"], [
        ["Εισαγωγή", "Τα δεδομένα είναι distributed και δεν θέλουμε central data collection.", "Μην ξεκινήσεις κατευθείαν από κώδικα."],
        ["FL/DFL", "FL κρατά raw data τοπικά. DFL αφαιρεί τον κεντρικό aggregation server.", "Μην πεις ότι FL από μόνο του λύνει όλη την ασφάλεια."],
        ["Baseline", "Πρώτα χτίστηκε καθαρό DFL για σημείο σύγκρισης.", "Μην ανακατέψεις baseline με security modes."],
        ["Masking", "Προστατεύει updates πριν την ανταλλαγή.", "Μην πεις ότι simple masking είναι πλήρες cryptographic protocol."],
        ["Pairwise masking", "Τα masks ακυρώνονται στο aggregation.", "Μην αγνοήσεις το floating point cancellation error, εξήγησέ το απλά."],
        ["Selective HE", "Ισχυρότερο privacy αλλά μεγαλύτερο overhead.", "Μην το παρουσιάσεις ως δωρεάν λύση."],
        ["Adaptive hybrid", "HE στους risky targets, masking στους υπόλοιπους.", "Μην πεις ότι βελτιώνει accuracy. Βελτιώνει trade-off."],
        ["Results", "Η accuracy μένει κοντά ίδια, ενώ το κόστος αλλάζει.", "Μην πανικοβληθείς αν οι curves επικαλύπτονται."],
        ["Blockchain", "Audit/integrity/tamper evidence.", "Μην το πεις production blockchain."],
        ["Platform demo", "Δείχνει πιο πρακτική λειτουργία με nodes/dashboard/verifier.", "Μην το μπερδέψεις με το core scientific experiment."],
    ], [1.45, 2.75, 2.3])

    doc.add_heading("21. Τι να μην πεις στην άμυνα", level=1)
    doc.add_paragraph(
        "Σε μία τεχνική παρουσίαση είναι καλύτερο να είσαι ακριβής παρά υπερβολικός. Τα παρακάτω είναι φράσεις που πρέπει να αποφύγεις ή να τις διορθώσεις."
    )
    add_table(doc, ["Μην πεις", "Πες καλύτερα"], [
        ["Το σύστημα είναι πλήρως ασφαλές.", "Το σύστημα ενσωματώνει και αξιολογεί μηχανισμούς προστασίας updates σε thesis prototype επίπεδο."],
        ["Το blockchain λύνει το privacy.", "Το blockchain προσφέρει integrity και auditability, όχι privacy."],
        ["Το adaptive hybrid βελτιώνει το accuracy.", "Το adaptive hybrid κρατά παρόμοιο accuracy με καλύτερο security/cost trade-off."],
        ["Το simple masking είναι production secure aggregation.", "Το simple masking είναι controlled simulation, ενώ το pairwise masking έχει ακύρωση masks."],
        ["Το CIFAR-10 έχει λυθεί πλήρως.", "Το CIFAR-10 υποστηρίζεται από το framework, αλλά για υψηλότερη απόδοση χρειάζεται CNN."],
        ["Το Docker demo είναι πλήρες distributed product.", "Το Docker demo είναι operational proof-of-concept για παρουσίαση και μελλοντική επέκταση."],
    ], [2.6, 3.9])

    doc.add_heading("22. Πώς να συνδέσεις κώδικα, Word και παρουσίαση", level=1)
    doc.add_paragraph(
        "Η τελική εικόνα που πρέπει να δώσεις είναι ότι τα τρία κομμάτια αλληλοσυμπληρώνονται."
    )
    add_table(doc, ["Artifact", "Ρόλος", "Τι αποδεικνύει"], [
        ["Κώδικας", "Υλοποίηση pipeline και security modes.", "Ότι η εργασία δεν είναι μόνο θεωρία."],
        ["CSV/plots", "Μετρήσιμα αποτελέσματα.", "Ότι τα συμπεράσματα βασίζονται σε πραγματικά runs."],
        ["Word αναφορά", "Αναλυτική τεχνική τεκμηρίωση.", "Ότι υπάρχει ακαδημαϊκή περιγραφή, μεθοδολογία και αποτελέσματα."],
        ["PowerPoint", "Προφορική παρουσίαση.", "Ότι μπορείς να εξηγήσεις τη συνεισφορά συνοπτικά."],
        ["Docker/platform demo", "Λειτουργική επίδειξη.", "Ότι υπάρχει κατεύθυνση προς πιο ρεαλιστικό σύστημα."],
    ], [1.65, 2.45, 2.4])

    doc.add_heading("23. Τι δεν ισχυριζόμαστε", level=1)
    add_bullets(doc, [
        "Δεν ισχυριζόμαστε ότι είναι production-ready blockchain.",
        "Δεν ισχυριζόμαστε ότι το simple masking είναι πλήρες secure aggregation.",
        "Δεν ισχυριζόμαστε ότι το HE είναι δωρεάν ή πάντα πρακτικό.",
        "Δεν ισχυριζόμαστε ότι το CIFAR-10 λύθηκε πλήρως με MLP.",
        "Δεν ισχυριζόμαστε ότι έχει γίνει πλήρες adversarial security evaluation.",
    ])
    add_callout(
        doc,
        "Γιατί αυτό είναι θετικό",
        "Σε διπλωματική εργασία είναι σημαντικό να ξέρεις τα όρια. Οι καθαροί περιορισμοί δείχνουν τεχνική ωριμότητα και όχι αδυναμία.",
    )

    doc.add_heading("24. Τι είναι η κύρια αφήγηση της παρουσίασης", level=1)
    add_numbers(doc, [
        "Ξεκινάς από το πρόβλημα: distributed data και privacy.",
        "Εξηγείς γιατί FL: τα δεδομένα μένουν τοπικά.",
        "Εξηγείς γιατί decentralized FL: δεν θέλουμε κεντρικό aggregation server.",
        "Δείχνεις baseline DFL.",
        "Προσθέτεις masking για προστασία updates.",
        "Προσθέτεις pairwise masking για σωστή ακύρωση masks.",
        "Προσθέτεις HE για ισχυρότερη προστασία αλλά δείχνεις ότι είναι ακριβό.",
        "Παρουσιάζεις adaptive hybrid ως trade-off.",
        "Δείχνεις results και λες ότι accuracy μένει σταθερή ενώ αλλάζει το κόστος.",
        "Κλείνεις με blockchain audit και platform demo ως ενίσχυση αξιοπιστίας/παρουσίασης.",
    ])

    doc.add_heading("25. Γρήγορο script 2 λεπτών", level=1)
    doc.add_paragraph(
        "Η εργασία μου ασχολείται με Secure Decentralized Federated Learning. Το βασικό πρόβλημα είναι ότι σε πολλά πραγματικά περιβάλλοντα "
        "τα δεδομένα είναι κατανεμημένα και δεν είναι επιθυμητό να μεταφερθούν σε έναν κεντρικό server. Για αυτό υλοποίησα ένα decentralized FL σύστημα, "
        "όπου κάθε client κρατά τα δεδομένα του τοπικά, εκπαιδεύει μοντέλο και ανταλλάσσει updates μόνο με γειτονικούς clients σύμφωνα με διαφορετικές topologies."
    )
    doc.add_paragraph(
        "Πάνω στο baseline πρόσθεσα μηχανισμούς προστασίας: simple masking, pairwise masking, selective homomorphic encryption και adaptive hybrid. "
        "Η κύρια συνεισφορά είναι το adaptive hybrid, γιατί εφαρμόζει HE στοχευμένα στους πιο risky clients και κρατά masking στους υπόλοιπους, ώστε να πετυχαίνει καλύτερο trade-off ανάμεσα σε privacy και κόστος."
    )
    doc.add_paragraph(
        "Επιπλέον πρόσθεσα blockchain-style audit layer και platform/Docker demo, ώστε να υπάρχει integrity verification, tamper evidence και πιο πρακτική παρουσίαση της λειτουργίας. "
        "Τα αποτελέσματα δείχνουν ότι οι security μηχανισμοί μπορούν να ενσωματωθούν χωρίς να χαλάσουν την accuracy, ενώ το κόστος διαφοροποιείται ανάλογα με το επίπεδο προστασίας."
    )

    doc.add_heading("26. Μεγάλο Q&A άμυνας", level=1)
    qa = [
        ("Ποιο είναι το αντικείμενο της διπλωματικής;",
         "Η υλοποίηση και αξιολόγηση ενός Secure Decentralized Federated Learning prototype με πολλαπλά security modes, topologies, datasets, metrics, plots και audit μηχανισμό."),
        ("Ποια είναι η βασική διαφορά FL και DFL;",
         "Στο κλασικό FL υπάρχει κεντρικός server για aggregation. Στο DFL οι clients επικοινωνούν peer-to-peer με βάση μία topology και κάνουν τοπικό aggregation."),
        ("Γιατί δεν στέλνουμε raw data;",
         "Για λόγους privacy, κανονισμών και πρακτικότητας. Στέλνουμε μόνο model updates, ώστε τα δεδομένα να παραμένουν στον client."),
        ("Το FL λύνει πλήρως το privacy πρόβλημα;",
         "Όχι. Μειώνει τη μεταφορά raw data, αλλά τα model updates μπορούν να διαρρεύσουν πληροφορία. Γι' αυτό προσθέτουμε masking και HE."),
        ("Γιατί χρειάζεται baseline;",
         "Για να ξέρουμε την απόδοση χωρίς security overhead και να συγκρίνουμε δίκαια κάθε μηχανισμό προστασίας."),
        ("Τι είναι topology;",
         "Είναι ο γράφος επικοινωνίας που ορίζει ποιοι clients ανταλλάσσουν updates μεταξύ τους."),
        ("Ποια topology είναι καλύτερη;",
         "Εξαρτάται. Fully connected διαχέει γρήγορα πληροφορία αλλά είναι ακριβό. Ring είναι φθηνό αλλά πιο αργό. Small-world είναι καλός συμβιβασμός."),
        ("Τι σημαίνει non-IID;",
         "Σημαίνει ότι οι clients έχουν διαφορετικές κατανομές δεδομένων, κάτι που κάνει το training πιο δύσκολο και πιο ρεαλιστικό."),
        ("Γιατί το non-IID έχει χαμηλότερη accuracy;",
         "Επειδή κάθε client εκπαιδεύεται σε biased subset, άρα τα local updates είναι πιο ανομοιογενή και το aggregation δυσκολεύεται."),
        ("Γιατί χρησιμοποιήθηκε MNIST;",
         "Είναι κλασικό, γρήγορο και σταθερό benchmark για να επιβεβαιωθεί ότι το pipeline δουλεύει σωστά."),
        ("Γιατί προστέθηκε Fashion-MNIST;",
         "Για να δοκιμαστεί το pipeline σε λίγο πιο δύσκολο dataset χωρίς μεγάλη αλλαγή input δομής."),
        ("Γιατί προστέθηκε CIFAR-10;",
         "Για να υπάρχει πιο δύσκολο RGB benchmark και να φανεί ότι το framework μπορεί να επεκταθεί πέρα από MNIST-like δεδομένα."),
        ("Γιατί το CIFAR-10 θέλει προσοχή;",
         "Επειδή το υπάρχον MLP δεν είναι ιδανικό για εικόνες RGB. Για υψηλότερη CIFAR απόδοση θα χρειαζόταν CNN."),
        ("Τι είναι state_dict;",
         "Είναι το σύνολο των παραμέτρων του μοντέλου στην PyTorch, δηλαδή weights και biases ανά layer."),
        ("Τι είναι aggregation;",
         "Ο συνδυασμός updates από γείτονες, συνήθως με averaging, ώστε να ενημερωθεί το τοπικό μοντέλο."),
        ("Τι κάνει το security_mode none;",
         "Τρέχει το καθαρό baseline χωρίς masking ή encryption."),
        ("Τι κάνει το simple masking;",
         "Προσθέτει τυχαία masks στα updates πριν διαμοιραστούν. Είναι controlled simulation protected sharing."),
        ("Είναι το simple masking πλήρες secure aggregation;",
         "Όχι. Είναι απλοποιημένη πρώτη έκδοση. Το pairwise masking είναι πιο σωστή προσέγγιση ακύρωσης masks."),
        ("Πώς δουλεύει το pairwise masking;",
         "Κάθε ζευγάρι clients χρησιμοποιεί masks με αντίθετα πρόσημα ώστε στο aggregation τα masks να ακυρώνονται."),
        ("Γιατί υπάρχει μικρό cancellation error;",
         "Λόγω floating point αριθμητικής. Όταν είναι τάξης 1e-7 δεν επηρεάζει ουσιαστικά το μοντέλο."),
        ("Τι είναι Homomorphic Encryption;",
         "Κρυπτογραφική τεχνική που επιτρέπει υπολογισμούς πάνω σε encrypted values."),
        ("Γιατί δεν χρησιμοποιούμε HE παντού;",
         "Επειδή έχει μεγάλο computational και communication overhead, ειδικά λόγω ciphertext expansion."),
        ("Τι είναι selective HE;",
         "Επιλεκτική εφαρμογή HE σε συγκεκριμένα updates ή παραμέτρους, αντί για πλήρη κρυπτογράφηση όλου του μοντέλου."),
        ("Ποια είναι η κύρια συνεισφορά;",
         "Το adaptive hybrid security mode, που συνδυάζει pairwise masking και HE με βάση risk score."),
        ("Τι είναι risk score;",
         "Μία μετρική που βοηθά να επιλεγούν οι πιο risky clients ή updates για HE προστασία."),
        ("Τι σημαίνει HE ratio 0.40;",
         "Περίπου 40% των clients/targets προστατεύονται με HE. Με 5 clients αυτό αντιστοιχεί σε περίπου 2 HE targets ανά round."),
        ("Γιατί η accuracy δεν αλλάζει πολύ ανά security mode;",
         "Επειδή οι μηχανισμοί προστασίας έχουν στόχο να προστατεύσουν την ανταλλαγή updates, όχι να αλλάξουν τη μαθησιακή διαδικασία."),
        ("Είναι καλό που η accuracy μένει ίδια;",
         "Ναι, γιατί δείχνει ότι το security layer δεν καταστρέφει το training."),
        ("Τι δείχνει το communication cost;",
         "Πόσα bytes χρειάζονται για τη διανομή updates, masks, ciphertexts και metadata."),
        ("Γιατί το HE αυξάνει το communication;",
         "Τα encrypted ciphertexts είναι πολύ μεγαλύτερα από τα απλά tensors."),
        ("Τι δείχνει το adaptive sensitivity analysis;",
         "Δείχνει πώς αλλάζει το κόστος όταν αυξάνεται το ποσοστό HE targets. Συνήθως η accuracy μένει σταθερή αλλά το κόστος αυξάνεται."),
        ("Τι ρόλο έχει το blockchain;",
         "Παρέχει audit trail, integrity verification και tamper evidence για τα communication rounds."),
        ("Το blockchain προστατεύει privacy;",
         "Όχι. Το blockchain layer είναι για integrity και auditability, όχι για απόκρυψη updates."),
        ("Είναι production blockchain;",
         "Όχι. Είναι lightweight educational/research blockchain component χωρίς real consensus ή validator network."),
        ("Τι είναι hash chain;",
         "Κάθε block περιέχει hash του προηγούμενου block. Έτσι αν αλλοιωθεί ένα block, σπάει η αλυσίδα."),
        ("Τι σημαίνει INVALID_HASH_CHAIN;",
         "Ότι το previous hash linking δεν συμφωνεί και πιθανόν έχει γίνει tampering ή αλλοίωση."),
        ("Τι σημαίνει INVALID_SIGNATURE;",
         "Ότι η υπογραφή δεν ταιριάζει με τα περιεχόμενα ή το client identity."),
        ("Τι προσφέρει το Docker demo;",
         "Δείχνει το σύστημα σαν μικρή distributed εφαρμογή με nodes, keys, audit logs, dashboard και verifier."),
        ("Το dashboard είναι επιστημονικό αποτέλεσμα;",
         "Όχι από μόνο του. Είναι εργαλείο παρακολούθησης και παρουσίασης του συστήματος."),
        ("Τι θα έκανες ως future work;",
         "CNN για CIFAR-10, μεγαλύτερα datasets, περισσότερους clients, adversarial attacks, real secure aggregation, real blockchain consensus και deployment σε πολλά machines."),
        ("Ποια είναι τα όρια της εργασίας;",
         "Είναι thesis prototype. Δεν είναι παραγωγικό security product και δεν έχει πλήρες adversarial evaluation."),
        ("Πώς απαντάς αν πουν ότι είναι simulation;",
         "Λέω ότι κάποια στοιχεία είναι πράγματι controlled simulation, αλλά υπάρχει πλήρες working pipeline, πραγματικά metrics, plots, Docker demo και σαφής διαχωρισμός περιορισμών."),
        ("Ποιο είναι το τελικό συμπέρασμα;",
         "Μπορεί να ενσωματωθεί προστασία σε DFL χωρίς μεγάλη απώλεια accuracy, αλλά υπάρχει trade-off σε χρόνο και communication. Το adaptive hybrid δίνει πιο ισορροπημένη λύση."),
        ("Γιατί πρόσθεσες experiment_name;",
         "Για να οργανώνονται τα αποτελέσματα ανά πείραμα και να μην γίνονται overwrite προηγούμενα CSV ή plots."),
        ("Γιατί χρειάζονται πολλά seeds;",
         "Για να μειωθεί η πιθανότητα ένα αποτέλεσμα να οφείλεται σε τυχαία αρχικοποίηση ή τυχαίο split."),
        ("Γιατί το GPU/CUDA έχει σημασία;",
         "Επιταχύνει το local training και κάνει πιο πρακτικό το τρέξιμο πολλών πειραμάτων, ειδικά όταν αυξάνονται rounds, clients ή datasets."),
        ("Το adaptive hybrid πώς αποφασίζει;",
         "Με βάση policy/risk score επιλέγει ένα ποσοστό targets για HE και αφήνει τα υπόλοιπα σε pairwise masking."),
        ("Τι σημαίνει overhead;",
         "Το επιπλέον κόστος που προσθέτει ένας μηχανισμός σε χρόνο, bytes ή υπολογιστικούς πόρους σε σχέση με το baseline."),
        ("Γιατί κρατήθηκε ίδια η training logic;",
         "Για να είναι δίκαιη η σύγκριση. Αν αλλάζαμε και το training, δεν θα ξέραμε αν οι διαφορές οφείλονται στο security mode ή στο μοντέλο/training."),
        ("Πώς ξέρεις ότι το pairwise masking δεν χαλάει το aggregation;",
         "Επειδή τα masks είναι σχεδιασμένα με αντίθετα πρόσημα και μετρήθηκε πολύ μικρό cancellation error."),
        ("Ποια είναι η διαφορά simulation και πραγματικής ασφάλειας;",
         "Simulation σημαίνει ότι μοντελοποιούμε τη συμπεριφορά για πειραματική αξιολόγηση. Πραγματική ασφάλεια απαιτεί αυστηρό πρωτόκολλο, threat model και adversarial validation."),
        ("Τι θα έβαζες αν είχες περισσότερο χρόνο;",
         "Πιο ισχυρό μοντέλο για CIFAR-10, real secure aggregation protocol, adversarial tests, deployment σε πολλά machines και production-grade key management."),
        ("Πώς θα το πουλούσες τεχνικά ως prototype;",
         "Ως privacy-aware decentralized learning and audit testbed για οργανισμούς που θέλουν να πειραματιστούν με distributed ML χωρίς κεντρική συγκέντρωση δεδομένων."),
    ]
    for q, a in qa:
        add_qa(doc, q, a)

    doc.add_heading("27. Πώς να χειριστείς δύσκολες ερωτήσεις", level=1)
    add_table(doc, ["Αν σε πιέσουν για...", "Απάντα με αυτή τη λογική"], [
        ["Production readiness", "Το παρουσιάζω ως research prototype, όχι ως production product."],
        ["Security guarantees", "Διαχωρίζω masking, pairwise masking, HE και audit. Δεν τα παρουσιάζω όλα ως ισοδύναμα."],
        ["CIFAR-10 accuracy", "Το framework υποστηρίζει CIFAR-10, αλλά το MLP δεν είναι ιδανικό. Future work: CNN."],
        ["Blockchain", "Είναι audit/integrity layer, όχι privacy layer και όχι production blockchain."],
        ["Overlapping plots", "Οι καμπύλες επικαλύπτονται επειδή οι μέθοδοι διατηρούν σχεδόν ίδια learning behavior."],
    ], [2.3, 4.2])

    doc.add_heading("28. Commands που πρέπει να ξέρεις", level=1)
    doc.add_paragraph("Δεν χρειάζεται να τα θυμάσαι όλα απ' έξω, αλλά πρέπει να ξέρεις τι αντιπροσωπεύουν.")
    commands = [
        "python src/main.py --security_mode none",
        "python src/main.py --security_mode pairwise_masking",
        "python src/main.py --security_mode selective_he",
        "python src/main.py --security_mode adaptive_hybrid",
        "python src/main.py --dataset fashion_mnist --topology star --security_mode none",
        "python src/main.py --dataset cifar10 --topology small_world --security_mode adaptive_hybrid",
        "python src/run_validation_experiments.py",
        "python src/run_adaptive_sensitivity.py --split_type non_iid",
        "python src/compare_results.py",
    ]
    for cmd in commands:
        add_code(doc, cmd)

    doc.add_heading("29. Τελικό checklist πριν την παρουσίαση", level=1)
    add_bullets(doc, [
        "Μπορώ να εξηγήσω τη διαφορά FL και DFL.",
        "Μπορώ να εξηγήσω IID και non-IID.",
        "Μπορώ να εξηγήσω κάθε topology με ένα πρακτικό trade-off.",
        "Μπορώ να εξηγήσω κάθε security mode χωρίς να υπερβάλλω.",
        "Μπορώ να πω γιατί το adaptive hybrid είναι η κύρια συνεισφορά.",
        "Μπορώ να εξηγήσω γιατί η accuracy παραμένει παρόμοια.",
        "Μπορώ να εξηγήσω γιατί το HE κοστίζει περισσότερο.",
        "Μπορώ να εξηγήσω ότι το blockchain είναι audit layer.",
        "Μπορώ να αναφέρω καθαρά τα όρια και future work.",
    ])

    doc.add_heading("30. Μία τελική καθαρή απάντηση αν σε ρωτήσουν τι έκανες", level=1)
    doc.add_paragraph(
        "Υλοποίησα ένα πλήρες thesis prototype για Secure Decentralized Federated Learning. Ξεκίνησα από baseline DFL, πρόσθεσα "
        "πολλαπλές topologies και datasets, υλοποίησα masking, pairwise masking, selective HE και adaptive hybrid security, κατέγραψα metrics και plots, "
        "έτρεξα συγκριτικά και validation experiments, πρόσθεσα blockchain-style audit για integrity και έφτιαξα Docker/platform demo με dashboard για πιο ρεαλιστική παρουσίαση. "
        "Η βασική τεχνική ιδέα είναι ότι η προστασία των updates έχει κόστος, άρα το adaptive hybrid προσπαθεί να χρησιμοποιήσει ακριβότερη προστασία μόνο εκεί που χρειάζεται περισσότερο."
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
