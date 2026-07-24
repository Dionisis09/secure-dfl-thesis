from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Secure_FL_Custom_RnD_Proposal.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def style_run(run, size=None, bold=False, color=None):
    run.font.name = "Calibri"
    if size:
        run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(5)
        run = p.add_run(text)
        style_run(run, size=14, bold=True, color="2E74B5")
    else:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(text)
        style_run(run, size=12, bold=True, color="1F4D78")
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.10
    run = p.add_run(text)
    style_run(run, size=10.5)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.10
    run = p.add_run(text)
    style_run(run, size=10)
    return p


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(3)
    r = title.add_run("Secure Federated Learning Custom R&D Pilot")
    style_run(r, size=20, bold=True, color="0B2545")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(8)
    r = subtitle.add_run("One-page proposal for a privacy-preserving ML proof-of-concept")
    style_run(r, size=11, color="555555")

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    labels = [
        ("Target investment", "10,000 EUR+"),
        ("Timeline", "4 weeks"),
        ("Output", "Prototype + report"),
    ]
    for idx, (label, value) in enumerate(labels):
        cell = table.rows[0].cells[idx]
        cell.width = Inches(2.15)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "F2F4F7")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        rr = p.add_run(label + "\n")
        style_run(rr, size=8.5, bold=True, color="1F4D78")
        rr = p.add_run(value)
        style_run(rr, size=10.5, bold=True, color="0B2545")

    add_heading(doc, "Executive summary")
    add_body(
        doc,
        "We propose a custom R&D pilot that evaluates whether privacy-preserving "
        "federated learning can solve a specific distributed-data machine learning "
        "problem without centralizing all raw data.",
    )

    add_heading(doc, "Client problem")
    add_body(
        doc,
        "Many organizations have useful data distributed across sites, departments, "
        "devices or partner institutions. Centralizing this data may be limited by "
        "privacy, ownership, compliance or operational constraints.",
    )

    add_heading(doc, "Proposed solution")
    add_body(
        doc,
        "We adapt a secure federated learning prototype to the client's use case. "
        "The system compares baseline learning against masking, selective homomorphic "
        "encryption, adaptive hybrid security and audit verification.",
    )

    add_heading(doc, "Deliverables")
    for item in [
        "Working prototype adapted to the client use case",
        "Baseline and secured federated learning experiments",
        "Metrics CSV files, plots and security overhead analysis",
        "Audit trail and tamper-detection demonstration",
        "Technical report, reproducibility commands and handover call",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Timeline")
    timeline = doc.add_table(rows=4, cols=2)
    timeline.alignment = WD_TABLE_ALIGNMENT.CENTER
    timeline.autofit = False
    weeks = [
        ("Week 1", "Requirements, data format, threat model and experiment design"),
        ("Week 2", "Prototype adaptation and baseline experiments"),
        ("Week 3", "Security modes, audit layer and comparison experiments"),
        ("Week 4", "Final analysis, report, handover and next-step roadmap"),
    ]
    for row, (week, detail) in zip(timeline.rows, weeks):
        row.cells[0].width = Inches(1.1)
        row.cells[1].width = Inches(5.0)
        for cell in row.cells:
            set_cell_margins(cell, top=60, bottom=60)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(row.cells[0], "E8EEF5")
        p = row.cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        rr = p.add_run(week)
        style_run(rr, size=9.5, bold=True, color="1F4D78")
        p = row.cells[1].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        rr = p.add_run(detail)
        style_run(rr, size=9.5)

    add_heading(doc, "Success criterion")
    add_body(
        doc,
        "The pilot is successful if the client receives a clear answer to whether "
        "secure federated learning is technically realistic for their use case, "
        "what trade-offs are expected and what should be built next.",
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(
        "Closing line: This pilot turns privacy-preserving machine learning from "
        "an abstract concept into a measurable, verifiable and decision-ready prototype."
    )
    style_run(r, size=10, bold=True, color="0B2545")

    doc.save(OUT)


if __name__ == "__main__":
    build()
