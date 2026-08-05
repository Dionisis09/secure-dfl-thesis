"""Append final Docker/pairwise-masking evidence to the technical report."""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def add_figure(doc: Document, image_path: Path, caption: str, width_in: float = 6.4) -> None:
    doc.add_paragraph()
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(width_in))
    caption_paragraph = doc.add_paragraph(caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption_paragraph.runs:
        run.font.size = Pt(9)
        run.italic = True


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    source = root / "docs" / "Secure_DFL_Technical_Report_COMPLETE.docx"
    output = root / "docs" / "Secure_DFL_Technical_Report_COMPLETE_WITH_PLATFORM_FINAL.docx"
    assets = root / "platform_mvp" / "results" / "presentation_assets"

    shutil.copy2(source, output)
    doc = Document(output)

    doc.add_page_break()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Appendix B - Final Platform Extension and Docker Real-Key Evidence")
    title_run.bold = True
    title_run.font.size = Pt(18)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Secure Decentralized Federated Learning Prototype")
    subtitle_run.italic = True
    subtitle_run.font.size = Pt(11)

    for text in [
        "In the final stage of the project, a more realistic platform-level demonstration was added. This extension does not modify the core thesis training pipeline. It improves the communication, control, auditability and presentation layer around the Secure DFL prototype.",
        "The new demo runs three independent Secure DFL nodes inside Docker, generates real Ed25519 deployment keys, uses a trusted public-key manifest, executes signed decentralized rounds and verifies the signed audit logs after execution.",
    ]:
        paragraph = doc.add_paragraph(text)
        paragraph.paragraph_format.space_after = Pt(6)

    doc.add_heading("B.1 Implemented platform additions", level=2)
    for item in [
        "Docker real-key demo with three independent Secure DFL nodes.",
        "Automatic generation of real Ed25519 deployment keys.",
        "Trusted public-key manifest for peer verification.",
        "Signed model-state envelopes and payload hash verification.",
        "New security mode: pairwise_masking with target-specific mask cancellation.",
        "Operator dashboard for node status, finalized rounds and masking metrics.",
        "Evidence export folder with demo summary, audit verification, Docker logs and per-node audit logs.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("B.2 Pairwise masking", level=2)
    for text in [
        "The new security mode is DFL_SECURITY_MODE=pairwise_masking. In this mode, raw masks are not transported inside the update envelope.",
        "Each contributor derives target-specific pairwise masks. During finalization, the target node applies its own local target-specific mask and then aggregates the masked contributor states. The pairwise masks cancel over the complete contributor set, so the final aggregate remains valid.",
        "This is stronger than the earlier controlled masking mode, where the mask was transported to the receiver for reconstruction.",
    ]:
        paragraph = doc.add_paragraph(text)
        paragraph.paragraph_format.space_after = Pt(6)

    doc.add_heading("B.3 Final validation and evidence", level=2)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Check"
    table.rows[0].cells[1].text = "Result"
    for left, right in [
        ("Platform validation matrix", "8/8 PASS"),
        ("Unit tests", "12/12 OK"),
        ("Docker real-key demo", "PASS"),
        ("Audit logs checked", "3"),
        ("Total audit records", "39"),
        ("Partial finalizations", "0"),
        ("Nodes online", "3/3"),
        ("Finalized rounds", "9"),
        ("Masked updates", "18"),
        ("Mask overhead", "1152 B"),
        ("Security mode", "pairwise_masking"),
    ]:
        cells = table.add_row().cells
        cells[0].text = left
        cells[1].text = right
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)

    evidence = doc.add_paragraph()
    evidence.paragraph_format.space_after = Pt(8)
    evidence.add_run("Evidence folder: ").bold = True
    evidence.add_run("platform_mvp/results/docker_real_key_demo_latest/")

    doc.add_heading("B.4 Figures from generated Docker results", level=2)
    add_figure(
        doc,
        assets / "dashboard_live.png",
        "Figure B.1 - Live operator dashboard. The dashboard shows 3/3 online nodes, 9 finalized rounds, 18 masked updates and pairwise masking.",
    )
    add_figure(
        doc,
        assets / "docker_evidence_summary.png",
        "Figure B.2 - Docker evidence summary generated from exported JSON artifacts. Audit verification is PASS for 3 node logs and 39 signed audit records.",
    )
    add_figure(
        doc,
        assets / "node0_status.png",
        "Figure B.3 - Node-level runtime status for an independent Secure DFL node, including finalized rounds, peer participation and runtime metrics.",
    )

    doc.add_heading("B.5 Limitations and next step", level=2)
    for item in [
        "The Docker demo uses local HTTP rather than production TLS or mTLS.",
        "The dashboard token is appropriate for demonstration, not enterprise authentication.",
        "The current pairwise_masking implementation requires the complete contributor set.",
        "Dropout-resilient pairwise masking and external audit anchoring remain future work.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    conclusion = doc.add_paragraph()
    conclusion.paragraph_format.space_before = Pt(8)
    conclusion.add_run("Conclusion. ").bold = True
    conclusion.add_run(
        "With this extension, the project includes both the research experiment layer and a practical platform-level demo with real deployment keys, signed communication, pairwise masked payload exchange, a live dashboard and verifiable audit logs."
    )

    doc.save(output)
    print(output)


if __name__ == "__main__":
    main()
